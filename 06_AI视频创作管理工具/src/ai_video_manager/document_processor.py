from __future__ import annotations

import re
from abc import ABC, abstractmethod
from dataclasses import dataclass, field
from pathlib import Path

from .models import Document, DocumentFormat, Script, Segment


@dataclass
class ScriptValidation:
    valid: bool
    missing_sections: list[str] = field(default_factory=list)
    reason: str = ""


class SplitStrategy(ABC):
    @abstractmethod
    def split(self, content: str) -> list[str]:
        raise NotImplementedError


class ChapterSplitStrategy(SplitStrategy):
    chinese_number_chars = r"零〇一二两三四五六七八九十百千万\d"
    chapter_pattern = re.compile(
        rf"""
        ^
        \s*
        (?:
            \#{{1,6}}\s+\S.* |
            第\s*[{chinese_number_chars}\s]+[章节回幕卷集部]\s*[:：、.\-\s]?.* |
            卷\s*[{chinese_number_chars}\s]+\s*[:：、.\-\s]?.* |
            (?:序章|楔子|尾声|后记|番外)(?:\s+\S.*)? |
            Chapter\s+(?:\d+|[A-Za-z]+|[IVXLCDM]+)\b.* |
            Part\s+(?:\d+|[A-Za-z]+|[IVXLCDM]+)\b.* |
            (?:\d+|[一二三四五六七八九十百千万]+)[、.]\s+\S.*
        )
        $
        """,
        re.IGNORECASE | re.MULTILINE | re.VERBOSE,
    )

    def split(self, content: str) -> list[str]:
        text = content.strip()
        if not text:
            return []

        matches = list(self.chapter_pattern.finditer(text))
        if not matches:
            return [text]

        segments: list[str] = []
        preamble = text[: matches[0].start()].strip()
        if preamble:
            segments.append(preamble)

        for index, match in enumerate(matches):
            start = match.start()
            end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
            segment = text[start:end].strip()
            if segment:
                segments.append(segment)
        return segments


class ManualMarkerSplitStrategy(SplitStrategy):
    marker_pattern = re.compile(r"(?m)^---\s*(?:segment|片段|分段)\s*---$")
    start_marker_pattern = re.compile(
        r"^\s*(?:【|\[)?\s*(?:第?\s*[一二三四五六七八九十百千万\d]+\s*段|片段\s*\d*|segment\s*\d*)\s*(?:始|开始|start)\s*(?:】|\])?\s*$",
        re.IGNORECASE,
    )
    end_marker_pattern = re.compile(
        r"^\s*(?:【|\[)?\s*(?:第?\s*[一二三四五六七八九十百千万\d]+\s*段|片段\s*\d*|segment\s*\d*)\s*(?:末|结束|end)\s*(?:】|\])?\s*$",
        re.IGNORECASE,
    )

    def split(self, content: str) -> list[str]:
        paired_segments = self._split_paired_markers(content)
        if paired_segments:
            return paired_segments
        parts = [part.strip() for part in self.marker_pattern.split(content)]
        return [part for part in parts if part]

    def _split_paired_markers(self, content: str) -> list[str]:
        segments: list[str] = []
        buffer: list[str] = []
        capturing = False
        saw_marker = False

        for line in content.splitlines():
            if self.start_marker_pattern.match(line):
                saw_marker = True
                buffer = []
                capturing = True
                continue
            if self.end_marker_pattern.match(line):
                saw_marker = True
                if capturing:
                    segment = "\n".join(buffer).strip()
                    if segment:
                        segments.append(segment)
                buffer = []
                capturing = False
                continue
            if capturing:
                buffer.append(line)

        if capturing:
            segment = "\n".join(buffer).strip()
            if segment:
                segments.append(segment)
        return segments if saw_marker else []


class CustomRegexSplitStrategy(SplitStrategy):
    regex_meta_chars = set(r"\.[](){}+^$|")

    def __init__(self, pattern: str) -> None:
        self.pattern = pattern.strip()
        if not self.pattern:
            raise ValueError("自定义切分符号不能为空")
        try:
            self.regex = re.compile(self._normalize_pattern(self.pattern), re.MULTILINE)
        except re.error as exc:
            raise ValueError(f"自定义切分正则无效：{exc}") from exc

    def split(self, content: str) -> list[str]:
        text = content.strip()
        if not text:
            return []

        matches = list(self.regex.finditer(text))
        matches = [match for match in matches if match.start() != match.end()]
        if not matches:
            return [text]

        segments: list[str] = []
        preamble = text[: matches[0].start()].strip()
        if preamble:
            segments.append(preamble)

        for index, match in enumerate(matches):
            start = match.start()
            end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
            segment = text[start:end].strip()
            if segment:
                segments.append(segment)
        return segments

    @classmethod
    def _normalize_pattern(cls, pattern: str) -> str:
        if ("*" in pattern or "?" in pattern) and not any(char in cls.regex_meta_chars for char in pattern):
            return re.escape(pattern).replace(r"\*", r".*").replace(r"\?", r".")
        return pattern


class LengthSplitStrategy(SplitStrategy):
    sentence_pattern = re.compile(r".+?(?:[。！？.!?]+|$)", re.DOTALL)

    def __init__(self, max_chars: int = 1200) -> None:
        if max_chars <= 0:
            raise ValueError("max_chars must be positive")
        self.max_chars = max_chars

    def split(self, content: str) -> list[str]:
        paragraphs = [p.strip() for p in re.split(r"\n\s*\n", content) if p.strip()]
        if not paragraphs:
            return []

        segments: list[str] = []
        buffer = ""
        for paragraph in paragraphs:
            for chunk in self._sentence_chunks(paragraph):
                candidate = f"{buffer}\n\n{chunk}".strip() if buffer else chunk
                if len(candidate) <= self.max_chars:
                    buffer = candidate
                    continue
                if buffer:
                    segments.append(buffer)
                buffer = chunk
        if buffer:
            segments.append(buffer)
        return segments

    def _sentence_chunks(self, paragraph: str) -> list[str]:
        sentences = [item.strip() for item in self.sentence_pattern.findall(paragraph) if item.strip()]
        chunks: list[str] = []
        buffer = ""
        for sentence in sentences:
            if len(sentence) > self.max_chars:
                if buffer:
                    chunks.append(buffer)
                    buffer = ""
                chunks.extend(sentence[i : i + self.max_chars] for i in range(0, len(sentence), self.max_chars))
                continue
            candidate = f"{buffer}{sentence}" if buffer else sentence
            if len(candidate) <= self.max_chars:
                buffer = candidate
                continue
            if buffer:
                chunks.append(buffer)
            buffer = sentence
        if buffer:
            chunks.append(buffer)
        return chunks


class DurationSplitStrategy(LengthSplitStrategy):
    def __init__(
        self,
        minutes: float = 2.0,
        chars_per_minute: int = 300,
        llm_client: object | None = None,
    ) -> None:
        super().__init__(max_chars=max(1, int(minutes * chars_per_minute)))
        self.minutes = minutes
        self.chars_per_minute = chars_per_minute
        self.llm_client = llm_client

    def split(self, content: str) -> list[str]:
        if not self.llm_client or not hasattr(self.llm_client, "suggest_duration_splits"):
            raise RuntimeError("Duration-based splitting requires a configured LLM client")
        segments = getattr(self.llm_client, "suggest_duration_splits")(content, self.minutes)
        if isinstance(segments, list) and all(isinstance(item, str) for item in segments):
            clean_segments = [segment.strip() for segment in segments if segment.strip()]
            if clean_segments:
                return clean_segments
        raise RuntimeError("LLM did not return valid duration-based segments")


class DocumentProcessor:
    supported_suffixes = {".txt", ".md", ".markdown", ".docx"}

    def load_document(self, file_path: str | Path, project_id: str) -> Document:
        path = Path(file_path)
        if path.suffix.lower() not in self.supported_suffixes:
            raise ValueError(f"Unsupported document format: {path.suffix}")

        if path.suffix.lower() == ".docx":
            content = self._load_docx(path)
            fmt = DocumentFormat.DOCX
        else:
            content = path.read_text(encoding="utf-8")
            fmt = self.detect_format(content, path)

        return Document(filename=path.name, content=content, format=fmt, project_id=project_id)

    def load_document_from_text(
        self,
        content: str,
        project_id: str,
        filename: str = "inline.txt",
    ) -> Document:
        return Document(
            filename=filename,
            content=content,
            format=self.detect_format(content, filename),
            project_id=project_id,
        )

    def detect_format(self, content: str, file_path: str | Path | None = None) -> DocumentFormat:
        suffix = Path(file_path).suffix.lower() if file_path else ""
        if suffix in {".md", ".markdown"}:
            return DocumentFormat.MARKDOWN
        if suffix == ".txt":
            return DocumentFormat.TXT
        if re.search(r"(?m)^#{1,6}\s+\S+", content) or re.search(r"\*\*[^*]+\*\*", content):
            return DocumentFormat.MARKDOWN
        return DocumentFormat.TXT if content.strip() else DocumentFormat.UNKNOWN

    def split_document(self, document: Document, strategy: SplitStrategy | None = None) -> list[Segment]:
        splitter = strategy or ChapterSplitStrategy()
        return [
            Segment(document_id=document.id, order=index + 1, content=content)
            for index, content in enumerate(splitter.split(document.content))
        ]

    def convert_to_script(self, segment: Segment) -> Script:
        content = segment.content.strip()
        lines = [line.strip() for line in content.splitlines() if line.strip()]
        title = lines[0].lstrip("#").strip() if lines else f"片段 {segment.order}"
        body = "\n".join(lines[1:] if len(lines) > 1 else lines)
        script = (
            f"【场景】{title}\n"
            f"【动作】镜头跟随文本节奏展开，保留原文的主要情节与情绪。\n"
            f"【对白】旁白：{body}"
        )
        return Script(segment_id=segment.id, content=script)

    def is_script_format(self, content: str) -> bool:
        text = content.strip()
        if not text:
            return False
        validation = self.validate_script(text)
        if validation.valid:
            return True
        lines = [line.strip() for line in text.splitlines() if line.strip()]
        if len(lines) < 3:
            return False
        scene_pattern = re.compile(
            r"^(?:【?场景】?|第[一二三四五六七八九十百千万\d]+场|[内外]景|INT\.|EXT\.)",
            re.IGNORECASE,
        )
        action_pattern = re.compile(r"^(?:【?动作】?|动作[:：]|镜头[:：]|画面[:：])")
        dialogue_pattern = re.compile(r"^(?:【?对白】?|[^：:\s]{1,16}[:：].+)")
        scene_count = sum(bool(scene_pattern.search(line)) for line in lines)
        action_count = sum(bool(action_pattern.search(line)) for line in lines)
        dialogue_count = sum(bool(dialogue_pattern.search(line)) for line in lines)
        return scene_count >= 1 and (action_count >= 1 or dialogue_count >= 2)

    def validate_script(self, script: str) -> ScriptValidation:
        text = script.strip()
        if not text:
            return ScriptValidation(valid=False, missing_sections=["内容"], reason="脚本为空")
        required = ["【场景】", "【动作】", "【对白】"]
        missing = [section for section in required if section not in text]
        if not missing:
            empty_sections = []
            for index, section in enumerate(required):
                start = text.find(section) + len(section)
                next_positions = [text.find(next_section, start) for next_section in required[index + 1 :]]
                next_positions = [position for position in next_positions if position >= 0]
                end = min(next_positions) if next_positions else len(text)
                if not text[start:end].strip():
                    empty_sections.append(section)
            if empty_sections:
                return ScriptValidation(
                    valid=False,
                    missing_sections=empty_sections,
                    reason=f"脚本结构不完整：{', '.join(empty_sections)}缺少内容",
                )
            return ScriptValidation(valid=True)

        lines = [line.strip() for line in text.splitlines() if line.strip()]
        scene_like = any(
            re.search(r"^(?:第[一二三四五六七八九十百千万\d]+场|[内外]景|INT\.|EXT\.)", line, re.IGNORECASE)
            for line in lines
        )
        dialogue_like = sum(bool(re.search(r"^[^：:\s]{1,16}[:：].+", line)) for line in lines)
        if scene_like and dialogue_like >= 2:
            return ScriptValidation(valid=True)
        return ScriptValidation(
            valid=not missing,
            missing_sections=missing,
            reason=f"缺少标准剧本结构：{', '.join(missing)}",
        )

    def suggest_split_points(
        self,
        content: str,
        llm_client: object | None = None,
        strategy_names: list[str] | None = None,
        max_chars: int = 1200,
        custom_pattern: str | None = None,
        duration_minutes: float = 2.0,
    ) -> list[dict[str, object]]:
        available: dict[str, SplitStrategy] = {
            "chapter": ChapterSplitStrategy(),
            "manual": ManualMarkerSplitStrategy(),
            "length": LengthSplitStrategy(max_chars=max_chars),
            "duration": DurationSplitStrategy(minutes=duration_minutes, llm_client=llm_client),
            "duration_2min": DurationSplitStrategy(minutes=duration_minutes, llm_client=llm_client),
        }
        if custom_pattern:
            available["custom_regex"] = CustomRegexSplitStrategy(custom_pattern)
        selected_names = strategy_names or ["chapter", "manual"]
        suggestions: list[dict[str, object]] = []
        for name in selected_names:
            if name not in available:
                raise ValueError(f"Unsupported split strategy: {name}")
            strategy = available[name]
            segments = strategy.split(content)
            suggestions.append(
                {
                    "strategy": name,
                    "segment_count": len(segments),
                    "summaries": [segment.splitlines()[0][:80] if segment else "" for segment in segments],
                    "estimated_minutes": round(sum(len(segment) for segment in segments) / 300, 1),
                }
            )
        return suggestions

    def _load_docx(self, path: Path) -> str:
        try:
            from docx import Document as DocxDocument  # type: ignore
        except ImportError as exc:
            raise RuntimeError("python-docx is required to load .docx files") from exc
        doc = DocxDocument(path)
        return "\n".join(paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip())
