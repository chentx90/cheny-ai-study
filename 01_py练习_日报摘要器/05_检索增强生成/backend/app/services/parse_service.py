"""解析服务：负责从上传文件生成 ParsedDoc 结构。

职责：文件加载 → 文本分块 → 标题/指纹计算 → LLM 富化（可选）。
无数据库依赖，全部逻辑无状态，使用模块级函数。
"""
import json
import re
from pathlib import Path
from typing import Optional


def _job_options(job: dict) -> dict:
    """JSONB 字段归一化为 dict。"""
    opts = job.get("options")
    if isinstance(opts, str):
        try:
            opts = json.loads(opts)
        except Exception:
            opts = {}
    return opts or {}


async def _resolve_doc_meta(text: str, job: dict, path: Path) -> tuple:
    """按优先级确定标题与指纹：手动 > LLM > 正文首行 > 文件名。"""
    from .doc_identity import compute_content_hash, derive_title
    from .llm_enrich import extract_title, is_available

    opts = _job_options(job)
    manual = job.get("title")
    domain = job.get("business_domain") or "general"
    llm_title = ""
    if not (manual and str(manual).strip()):
        if (opts.get("generate_summary") or opts.get("generate_questions") or opts.get("extract_title")) and is_available():
            llm_title = await extract_title(text)
    return derive_title(text, manual=manual, llm_title=llm_title, filename_stem=path.stem), compute_content_hash(text, domain)


async def _build_text_chunks(text: str, title: str, options: dict) -> list:
    """文本 → 分块列表，按 options 开关做 LLM 富化。"""
    from ..core.chunker import Chunker
    from .llm_enrich import enrich_chunks

    chunker = Chunker(
        title_context=title,
        section_path=f"{title} > 全文",
        max_chars=options.get("chunk_size", 600),
        overlap=options.get("chunk_overlap", 80),
    )
    chunks = [
        {
            "chunk_index": idx,
            "section_path": c.section_path,
            "title_context": c.title_context,
            "content": c.content,
            "summary": c.summary,
            "preset_questions": c.preset_questions,
            "physical_context": {"page": 1},
        }
        for idx, c in enumerate(chunker.chunk_text(text))
    ]
    await enrich_chunks(
        chunks,
        gen_summary=bool(options.get("generate_summary")),
        gen_questions=bool(options.get("generate_questions")),
    )
    return chunks


async def _build_page_chunks(
    pages: list[dict],
    title: str,
    options: dict,
    heading_label: str = "Page",
    physical_key: str = "page",
) -> list:
    """Build chunks page by page so retrieval can point back to the source page."""
    from ..core.chunker import Chunker
    from .llm_enrich import enrich_chunks

    chunks = []
    for page in pages:
        page_number = page.get("page_number") or 1
        heading_pattern = rf'^##\s+{re.escape(heading_label)}\s+\d+\s*'
        page_text = re.sub(heading_pattern, '', page.get("text", ""), count=1).strip()
        chunker = Chunker(
            title_context=title,
            section_path=f"{title} > {heading_label} {page_number}",
            max_chars=options.get("chunk_size", 600),
            overlap=options.get("chunk_overlap", 80),
        )
        for c in chunker.chunk_text(page_text):
            chunks.append({
                "chunk_index": len(chunks),
                "section_path": c.section_path,
                "title_context": c.title_context,
                "content": c.content,
                "summary": c.summary,
                "preset_questions": c.preset_questions,
                "physical_context": {physical_key: page_number},
            })
    await enrich_chunks(
        chunks,
        gen_summary=bool(options.get("generate_summary")),
        gen_questions=bool(options.get("generate_questions")),
    )
    return chunks


async def _parse_txt(path: Path, job: dict) -> dict:
    text = path.read_text(encoding="utf-8", errors="replace")
    options = _job_options(job)
    title, content_hash = await _resolve_doc_meta(text, job, path)
    chunks = await _build_text_chunks(text, title, options)
    return {
        "document": {"title": title, "doc_type": job.get("doc_type", "text"),
                     "business_domain": job.get("business_domain", "general"),
                     "source_uri": str(path), "content_hash": content_hash},
        "pages": [{"page_number": 1, "text": text}], "chunks": chunks,
        "data_assets": [], "sheets": [], "errors": [],
        "stats": {"pages": 1, "chunks": len(chunks), "data_assets": 0, "sheets": 0, "errors": 0},
    }


async def _parse_docx(path: Path, job: dict) -> dict:
    from .docx_extractor import extract_docx_markdown

    extracted = extract_docx_markdown(path)
    text = extracted.text
    options = _job_options(job)
    title, content_hash = await _resolve_doc_meta(text, job, path)
    chunks = await _build_text_chunks(text, title, options)
    return {
        "document": {"title": title, "doc_type": job.get("doc_type", "document"),
                     "business_domain": job.get("business_domain", "general"),
                     "source_uri": str(path), "content_hash": content_hash},
        "pages": [{"page_number": 1, "text": text, "metadata": {
            "paragraphs": extracted.paragraphs,
            "tables": extracted.tables,
        }}], "chunks": chunks,
        "data_assets": [], "sheets": [],
        "errors": [],
        "stats": {"pages": 1, "chunks": len(chunks), "data_assets": 0,
                  "sheets": extracted.tables, "errors": 0},
    }


async def _parse_pptx(path: Path, job: dict) -> dict:
    from pptx import Presentation
    prs = Presentation(path)
    pages = []
    for i, slide in enumerate(prs.slides, 1):
        texts = _extract_slide_texts(slide.shapes)
        if texts:
            pages.append({"page_number": i, "text": f"## Slide {i}\n\n" + "\n".join(texts)})
    text = "\n\n".join(page["text"] for page in pages)
    options = _job_options(job)
    title, content_hash = await _resolve_doc_meta(text, job, path)
    chunks = await _build_page_chunks(
        pages,
        title,
        options,
        heading_label="Slide",
        physical_key="slide",
    )
    return {
        "document": {"title": title, "doc_type": job.get("doc_type", "presentation"),
                     "business_domain": job.get("business_domain", "general"),
                     "source_uri": str(path), "content_hash": content_hash},
        "pages": pages,
        "chunks": chunks, "data_assets": [], "sheets": [], "errors": [],
        "stats": {"pages": len(pages), "chunks": len(chunks), "data_assets": 0, "sheets": 0, "errors": 0},
    }


def _extract_slide_texts(shapes) -> list[str]:
    texts: list[str] = []
    for shape in shapes:
        if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
            texts.append(_clean_pptx_text(shape.text_frame.text))
        if getattr(shape, "has_table", False):
            rows = []
            for row in shape.table.rows:
                cells = [_clean_pptx_text(cell.text).replace("\n", " ") for cell in row.cells]
                if any(cells):
                    rows.append(" | ".join(cells))
            if rows:
                texts.append("\n".join(rows))
        if hasattr(shape, "shapes"):
            texts.extend(_extract_slide_texts(shape.shapes))
    return texts


def _clean_pptx_text(text: str) -> str:
    return re.sub(r'\n{3,}', '\n\n', text.replace("\x0b", "\n")).strip()


async def parse(upload: dict, job: dict) -> dict:
    """统一解析入口，按 source_type 分发。"""
    source_type = job["source_type"]
    # 路径相对于 backend/ 根目录
    disk_path = Path(__file__).parent.parent.parent / upload["storage_uri"]
    if not disk_path.exists():
        raise FileNotFoundError(f"Uploaded file not found: {disk_path}")

    handlers = {
        "pdf": _parse_pdf,
        "excel": _parse_excel,
        "markdown": _parse_markdown,
        "json": _parse_json,
        "txt": _parse_txt,
        "docx": _parse_docx,
        "pptx": _parse_pptx,
    }
    handler = handlers.get(source_type)
    if not handler:
        raise ValueError(f"Unsupported source type: {source_type}")
    return await handler(disk_path, job)


async def _parse_pdf(path: Path, job: dict) -> dict:
    import fitz
    from .ocr_client import ocr_pdf_page
    options = _job_options(job)
    doc = fitz.open(path)
    pages = []
    for i in range(len(doc)):
        page = doc.load_page(i)
        text = page.get_text("text").strip()
        # OCR 调用条件：扫描页（无文字） 或 图片嵌入为主（文字极少）
        if not text or (page.get_images() and len(text) < 50):
            ocr_text = await ocr_pdf_page(page, i)
            if ocr_text:
                text = ocr_text
        if text:
            pages.append({"page_number": i + 1, "text": f"## Page {i + 1}\n\n{text}"})
    doc.close()
    full_text = "\n\n".join(p["text"] for p in pages)
    title, content_hash = await _resolve_doc_meta(full_text, job, path)
    chunks = await _build_page_chunks(pages, title, options)
    return {
        "document": {"title": title, "doc_type": job.get("doc_type", "manual"),
                     "business_domain": job.get("business_domain", "general"), "source_uri": str(path), "content_hash": content_hash},
        "pages": pages, "chunks": chunks, "data_assets": [], "sheets": [], "errors": [],
        "stats": {"pages": len(pages), "chunks": len(chunks), "data_assets": 0, "sheets": 0, "errors": 0},
    }


async def _parse_markdown(path: Path, job: dict) -> dict:
    text = path.read_text(encoding="utf-8")
    options = _job_options(job)
    title, content_hash = await _resolve_doc_meta(text, job, path)
    chunks = await _build_text_chunks(text, title, options)
    return {
        "document": {"title": title, "doc_type": job.get("doc_type", "markdown"),
                     "business_domain": job.get("business_domain", "general"),
                     "source_uri": str(path), "content_hash": content_hash},
        "pages": [{"page_number": 1, "text": text}], "chunks": chunks,
        "data_assets": [], "sheets": [], "errors": [],
        "stats": {"pages": 1, "chunks": len(chunks), "data_assets": 0, "sheets": 0, "errors": 0},
    }


async def _parse_json(path: Path, job: dict) -> dict:
    data = json.loads(path.read_text(encoding="utf-8"))
    title = job.get("title") or path.stem
    if isinstance(data, list):  # data_asset 列表场景
        return {
            "document": {"title": title, "doc_type": job.get("doc_type", "json"),
                         "business_domain": job.get("business_domain", "general"), "source_uri": str(path)},
            "pages": [], "chunks": [], "data_assets": data, "sheets": [], "errors": [],
            "stats": {"pages": 0, "chunks": 0, "data_assets": len(data), "sheets": 0, "errors": 0},
        }
    text = json.dumps(data, ensure_ascii=False, indent=2)
    options = _job_options(job)
    title, content_hash = await _resolve_doc_meta(text, job, path)
    chunks = await _build_text_chunks(text, title, options)
    return {
        "document": {"title": title, "doc_type": job.get("doc_type", "json"),
                     "business_domain": job.get("business_domain", "general"),
                     "source_uri": str(path), "content_hash": content_hash},
        "pages": [{"page_number": 1, "text": text}], "chunks": chunks,
        "data_assets": [], "sheets": [], "errors": [],
        "stats": {"pages": 1, "chunks": len(chunks), "data_assets": 0, "sheets": 0, "errors": 0},
    }


async def _parse_excel(path: Path, job: dict) -> dict:
    import pandas as pd
    options = _job_options(job)
    if options.get("excel_mode") == "text":
        return await _parse_excel_as_text(path, job, options)
    excel_file = pd.ExcelFile(path)
    sheets, data_assets = [], []
    for sheet_name in excel_file.sheet_names:
        if sheet_name not in {"tables", "columns", "metrics", "cases", "tools"}:
            continue
        df = pd.read_excel(excel_file, sheet_name=sheet_name).fillna("")
        rows = [{"asset_type": sheet_name.rstrip("s"), **{c: str(row[c]).strip() for c in df.columns}}
                for _, row in df.iterrows()]
        sheets.append({"sheet_name": sheet_name, "rows": len(df), "data_assets": rows})
        data_assets.extend(rows)
    return {
        "document": None, "pages": [], "chunks": [], "data_assets": data_assets,
        "sheets": sheets, "errors": [],
        "stats": {"pages": 0, "chunks": 0, "data_assets": len(data_assets), "sheets": len(sheets), "errors": 0},
    }


async def _parse_excel_as_text(path: Path, job: dict, options: dict) -> dict:
    from .excel_extractor import extract_excel_markdown

    extracted = extract_excel_markdown(
        path,
        max_table_cols=options.get("excel_max_table_cols", 18),
        anchor_cols=options.get("excel_anchor_cols", 2),
    )
    full_text = extracted.text
    title, content_hash = await _resolve_doc_meta(full_text, job, path)
    chunks = await _build_text_chunks(full_text, title, options)
    return {
        "document": {"title": title, "doc_type": job.get("doc_type", "excel"),
                     "business_domain": job.get("business_domain", "general"),
                     "source_uri": str(path), "content_hash": content_hash},
        "pages": [
            {"page_number": i + 1, "text": sheet.text, "metadata": {
                "sheet": sheet.name,
                "rows": sheet.rows,
                "columns": sheet.columns,
                "tables": sheet.tables,
            }}
            for i, sheet in enumerate(extracted.sheets)
        ],
        "chunks": chunks,
        "data_assets": [],
        "sheets": [
            {"sheet_name": sheet.name, "rows": sheet.rows, "columns": sheet.columns, "tables": sheet.tables}
            for sheet in extracted.sheets
        ],
        "errors": [],
        "stats": {"pages": len(extracted.sheets), "chunks": len(chunks), "data_assets": 0,
                  "sheets": len(extracted.sheets), "errors": 0},
    }
