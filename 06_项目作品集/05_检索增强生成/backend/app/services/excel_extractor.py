"""Excel to Markdown extraction helpers."""
from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


@dataclass
class ExtractedExcelSheet:
    name: str
    text: str
    rows: int
    columns: int
    tables: int


@dataclass
class ExtractedExcel:
    text: str
    sheets: list[ExtractedExcelSheet]


def extract_excel_markdown(
    file_path: str | Path,
    max_table_cols: int = 18,
    anchor_cols: int = 2,
) -> ExtractedExcel:
    path = Path(file_path)
    values_wb = load_workbook(path, data_only=True, read_only=False)
    formulas_wb = load_workbook(path, data_only=False, read_only=False)
    sheets: list[ExtractedExcelSheet] = []

    for values_ws in values_wb.worksheets:
        formulas_ws = formulas_wb[values_ws.title]
        grid, min_row, max_row, used_cols = _read_used_grid(values_ws, formulas_ws)
        if not grid or not used_cols:
            continue

        parts = [f"## {values_ws.title}"]
        table_count = 0
        for cols in _column_windows(used_cols, max_table_cols=max_table_cols, anchor_cols=anchor_cols):
            start = get_column_letter(cols[0])
            end = get_column_letter(cols[-1])
            parts.append(f"### {values_ws.title}!{start}{min_row}:{end}{max_row}")
            parts.append(_render_markdown_table(grid, cols))
            table_count += 1

        sheets.append(ExtractedExcelSheet(
            name=values_ws.title,
            text="\n\n".join(parts),
            rows=max_row - min_row + 1,
            columns=len(used_cols),
            tables=table_count,
        ))

    return ExtractedExcel(
        text="\n\n".join(sheet.text for sheet in sheets),
        sheets=sheets,
    )


def _read_used_grid(values_ws, formulas_ws):
    vertical_merges = _vertical_merge_lookup(values_ws)
    rows = []
    used_cols = set()
    min_row = None
    max_row = 0

    for row_idx in range(1, values_ws.max_row + 1):
        row_values = {}
        has_value = False
        for col_idx in range(1, values_ws.max_column + 1):
            value = _cell_value(values_ws, formulas_ws, row_idx, col_idx, vertical_merges)
            text = _format_value(value)
            if text:
                row_values[col_idx] = text
                used_cols.add(col_idx)
                has_value = True
        if has_value:
            min_row = row_idx if min_row is None else min_row
            max_row = row_idx
        rows.append(row_values)

    if min_row is None:
        return [], 0, 0, []

    return rows[min_row - 1:max_row], min_row, max_row, sorted(used_cols)


def _vertical_merge_lookup(ws) -> dict[tuple[int, int], Any]:
    lookup = {}
    for merged in ws.merged_cells.ranges:
        if merged.min_col == merged.max_col and merged.min_row < merged.max_row:
            value = ws.cell(merged.min_row, merged.min_col).value
            for row_idx in range(merged.min_row + 1, merged.max_row + 1):
                lookup[(row_idx, merged.min_col)] = value
    return lookup


def _cell_value(values_ws, formulas_ws, row_idx: int, col_idx: int, vertical_merges: dict):
    value = values_ws.cell(row_idx, col_idx).value
    if value is None:
        value = vertical_merges.get((row_idx, col_idx))
    if value is None:
        formula = formulas_ws.cell(row_idx, col_idx).value
        if isinstance(formula, str) and formula.startswith("="):
            value = formula
    return value


def _column_windows(used_cols: list[int], max_table_cols: int, anchor_cols: int) -> list[list[int]]:
    if len(used_cols) <= max_table_cols:
        return [used_cols]

    anchors = used_cols[:min(anchor_cols, len(used_cols))]
    payload_width = max(1, max_table_cols - len(anchors))
    payload = used_cols[len(anchors):]
    windows = []
    for start in range(0, len(payload), payload_width):
        windows.append(anchors + payload[start:start + payload_width])
    return windows


def _render_markdown_table(grid: list[dict[int, str]], cols: list[int]) -> str:
    header = ["Row"] + [get_column_letter(col) for col in cols]
    lines = [_format_row(header), _format_row(["---"] * len(header))]
    for offset, row_values in enumerate(grid, 1):
        row = [str(offset)] + [row_values.get(col, "") for col in cols]
        lines.append(_format_row(row))
    return "\n".join(lines)


def _format_row(values: list[str]) -> str:
    return "| " + " | ".join(_escape_cell(value) for value in values) + " |"


def _escape_cell(value: str) -> str:
    return str(value).replace("|", "\\|").replace("\n", "<br>").strip()


def _format_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return f"{value:.6g}"
    return str(value).strip()
